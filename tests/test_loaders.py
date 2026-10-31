import pytest
from docx import Document

from cv_analyser.exceptions import (
    CorruptedFileError,
    EmptyDocumentError,
    UnsupportedFileTypeError,
)
from cv_analyser.loaders.dispatcher import load_document


def test_unsupported_extension(tmp_path):
    file = tmp_path / "cv.txt"
    file.write_text("hello")
    with pytest.raises(UnsupportedFileTypeError):
        load_document(file)


def test_corrupted_pdf(tmp_path):
    file = tmp_path / "bad.pdf"
    file.write_bytes(b"this is not a pdf")
    with pytest.raises(CorruptedFileError):
        load_document(file)


def test_corrupted_docx(tmp_path):
    file = tmp_path / "bad.docx"
    file.write_bytes(b"this is not a docx")
    with pytest.raises(CorruptedFileError):
        load_document(file)


def test_empty_docx(tmp_path):
    file = tmp_path / "empty.docx"
    Document().save(file)
    with pytest.raises(EmptyDocumentError):
        load_document(file)


def test_docx_reads_paragraphs_and_tables(tmp_path):
    file = tmp_path / "cv.docx"
    doc = Document()
    doc.add_paragraph("Rahim Uddin")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Python"
    table.rows[0].cells[1].text = "SQL"
    doc.save(file)

    text = load_document(file)
    assert "Rahim Uddin" in text
    assert "Python" in text and "SQL" in text