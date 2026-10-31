from pathlib import Path
from zipfile import BadZipFile

from docx import Document
from docx.opc.exceptions import PackageNotFoundError

from cv_analyser.exceptions import CorruptedFileError, EmptyDocumentError


def load_docx(path: Path) -> str:
    try:
        document = Document(path)
    except (PackageNotFoundError, BadZipFile) as exc:
        raise CorruptedFileError(f"Cannot read DOCX: {path.name}") from exc

    lines = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            lines.extend(cell.text for cell in row.cells)

    text = "\n".join(lines).strip()
    if not text:
        raise EmptyDocumentError(f"No extractable text in DOCX: {path.name}")
    return text
